"""Focused offline v12 checks: corrective project portfolio and Good Point compliance."""
from playwright.sync_api import sync_playwright, expect
from openpyxl import load_workbook
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile
import argparse,json,os,re
p=argparse.ArgumentParser();p.add_argument('--url',default='http://127.0.0.1:8772/Didban_Amniat_Offline.html');p.add_argument('--artifacts',default='/tmp/didban-v12-qa');args=p.parse_args();art=Path(args.artifacts);art.mkdir(parents=True,exist_ok=True)
with TemporaryDirectory() as tmp,sync_playwright() as pw:
 out=Path(tmp);b=pw.chromium.launch(executable_path=os.environ.get('DIDBAN_CHROMIUM','/usr/bin/chromium'),headless=True,args=['--no-sandbox','--disable-dev-shm-usage']);ctx=b.new_context(viewport={'width':1536,'height':1050},accept_downloads=True);page=ctx.new_page();errors=[];requests=[];page.on('pageerror',lambda e:errors.append(str(e)));page.on('request',lambda r:requests.append(r.url));page.goto(args.url)
 def main(name):page.locator('aside nav>button').filter(has_text=name).click()
 def audit(value,reviewer='ممیز آزمون'):
  page.get_by_label('نام ممیز',exact=True).fill(reviewer);field=page.get_by_label('انطباق با استاندارد',exact=True);assert field.locator('option').count()==7;field.select_option(str(value));return field
 def saveaudit():page.get_by_role('button',name='ذخیره ارزیابی',exact=True).click();expect(page.get_by_role('dialog')).to_have_count(0)
 # The new result is independent of implementation and round-trips through native audit storage.
 page.locator('[data-standard-dot="A.5.1"]').click();field=audit(5);assert 'Good Point' in field.inner_text();assert field.evaluate('e=>getComputedStyle(e).backgroundColor')=='rgb(126, 34, 206)'
 page.get_by_text('اطلاعات ممیزی قبلی (در صورت وجود)',exact=True).click();prior=page.get_by_label('وضعیت انطباق در ممیزی قبلی',exact=True);assert prior.locator('option').count()==7;prior.select_option('5');assert prior.evaluate('e=>getComputedStyle(e).backgroundColor')=='rgb(126, 34, 206)';saveaudit()
 page.locator('[data-standard-dot="A.5.2"]').click();audit(2);saveaudit();main('الزامات استاندارد');page.get_by_role('cell',name='4.1',exact=True).click();audit(5);saveaudit();page.reload();main('داشبورد مدیریتی')
 dot=page.locator('[data-standard-dot="A.5.1"]');assert dot.get_attribute('data-compliance')=='5';assert dot.get_attribute('data-compliance-rank')=='5';assert page.locator('[data-standard-dot]').count()==93
 summary=page.locator('.dc-summary article').filter(has_text='Good Point');expect(summary.locator('strong')).to_have_text('۲');assert page.locator('[data-compliance-column="5"]').count()>1;assert page.locator('[data-compliance-series="5"]').count()>0
 assert page.locator('.dc-section thead [data-compliance-column="5"]').inner_text()=='نقطه قوت (Good Point)';page.locator('.dc-section').screenshot(path=str(art/'goodpoint-summary-light.png'))
 print('PASS current/prior Good Point, clauses+controls save/reload, purple radar layer, summary/chart/table counts',flush=True)
 # Project creation follows the real form, with high-level scope selection separate from exact links.
 main('اقدامات اصلاحی')
 projects=[('پروژه حاکمیت','1','active','25',['شناخت سازمان','حاکمیت و انطباق']),('پروژه افراد','2','blocked','60',['افراد و دورکاری']),('پروژه رهبری','3','completed','100',['رهبری'])]
 for i,(title,months,status,progress,scopes) in enumerate(projects):
  page.get_by_role('button',name='تعریف اقدام بهبود',exact=True).click();form=page.get_by_role('dialog');expect(form).to_be_visible()
  for label,value in [('عنوان اقدام یا پروژه',title),('مسئول پیگیری','مالک '+title),('هدف و نتیجه مورد انتظار','کاهش شکاف استاندارد'),('منابع لازم','تیم اجرا')]:form.get_by_label(label,exact=True).fill(value)
  form.get_by_label('گروه زمانی پروژه',exact=True).select_option(months);form.get_by_label('وضعیت پروژه',exact=True).select_option(status);
  if status=='completed':expect(form.get_by_label('پیشرفت (درصد)',exact=True)).to_have_value('100')
  else:form.get_by_label('پیشرفت (درصد)',exact=True).fill(progress)
  form.get_by_label('تاریخ شروع پروژه',exact=True).fill('۱۴۰۵/۰۸/۰۱')
  for heading in scopes:form.get_by_label('حوزه پروژه '+heading,exact=True).check()
  if i==1:
   form.locator('summary').filter(has_text='ارتباط دقیق با الزامات و کنترل‌ها').click();form.get_by_label('ارتباط پروژه A.5.1',exact=True).check()
  if i==0:
   form.locator('summary').filter(has_text='بودجه اقدام اصلاحی').click();form.get_by_label('بودجه مصوب اقدام اصلاحی',exact=True).fill('1000');form.get_by_label('تاریخ تصویب بودجه',exact=True).fill('۱۴۰۵/۰۸/۰۱')
   for j,(amount,date) in enumerate([('100','۱۴۰۵/۰۸/۰۲'),('250','۱۴۰۵/۰۸/۰۳')],1):
    form.get_by_role('button',name='افزودن ردیف تخصیص بودجه',exact=True).click();form.get_by_label('میزان بودجه تخصیص داده شده '+str(j),exact=True).fill(amount);form.get_by_label('تاریخ تخصیص بودجه '+str(j),exact=True).fill(date)
   assert '۳۵۰' in form.locator('[data-v12-budget-total]').inner_text();assert '۶۵۰' in form.locator('[data-v12-budget-remaining]').inner_text();form.screenshot(path=str(art/'project-budget-form.png'))
  form.get_by_role('button',name='ذخیره پروژه',exact=True).click();expect(form).to_have_count(0)
 assert page.locator('[data-improvement-project]').count()==3
 colors=[]
 for group in ['A','B','C']:
  card=page.locator('[data-improvement-project][data-project-group="'+group+'"]');assert card.count()==1;colors.append(card.evaluate('e=>getComputedStyle(e).borderTopColor'))
 assert len(set(colors))==3,colors
 # Every axis and progress value is computed from stored projects, without deduplicating distinct scopes.
 assert page.locator('[data-v12-scope-radar=clause]').count()==1
 assert page.locator('[data-v12-scope="clause:4"]').get_attribute('data-count')=='1'
 assert page.locator('[data-v12-scope="clause:5"]').get_attribute('data-count')=='1'
 page.get_by_label('دامنه نمودار گستره پروژه‌ها',exact=True).select_option('control')
 assert page.locator('[data-v12-scope-radar=control]').count()==1
 assert page.locator('[data-v12-scope="annex:governance"]').get_attribute('data-count')=='2'
 assert page.locator('[data-v12-scope="annex:people"]').get_attribute('data-count')=='1'
 panel=page.locator('[data-v12-project-dashboard=corrective]');assert panel.locator('[data-v12-progress]').count()==3
 assert '۶۲٪' in panel.locator('.v12-average').inner_text()
 for group,progress in [('A','۲۵'),('B','۶۰'),('C','۱۰۰')]:
  groupcard=panel.locator('[data-v12-project-group-summary="'+group+'"]');assert '۱ پروژه' in groupcard.inner_text();assert progress+'٪' in groupcard.inner_text()
 page.locator('.v12-project-charts').screenshot(path=str(art/'corrective-dashboard-light.png'))
 page.get_by_label('فیلتر گروه زمانی پروژه',exact=True).select_option('B');assert page.locator('[data-improvement-project]').count()==1;page.get_by_label('فیلتر گروه زمانی پروژه',exact=True).select_option('all')
 main('داشبورد مدیریتی');assert page.locator('[data-v12-project-dashboard=main] [data-v12-progress]').count()==3;main('اقدامات اصلاحی')
 page.locator('[data-improvement-project]').first.screenshot(path=str(art/'project-card.png'))
 page.get_by_role('button',name='ویرایش پروژه پروژه حاکمیت',exact=True).click();form=page.get_by_role('dialog');assert form.get_by_label('حوزه پروژه شناخت سازمان',exact=True).is_checked();assert form.get_by_label('حوزه پروژه حاکمیت و انطباق',exact=True).is_checked();form.locator('summary').filter(has_text='بودجه اقدام اصلاحی').click();expect(form.get_by_label('بودجه مصوب اقدام اصلاحی',exact=True)).to_have_value('1000');assert form.get_by_label('تاریخ تخصیص بودجه 2',exact=True).input_value()=='۱۴۰۵/۰۸/۰۳';form.get_by_role('button',name='ذخیره پروژه',exact=True).click()
 print('PASS scope-only and exact-link project creation, A/B/C cards, repeated allocations350/remain650 and editable persisted budget dates',flush=True)
 # Portable project backup includes both additions; restore into another workspace exercises the native validators.
 main('پشتیبانی سایت')
 with page.expect_download() as d:page.get_by_role('button',name='دریافت JSON پروژه فعال',exact=True).click()
 d.value.save_as(str(out/'backup.json'));backup=json.loads((out/'backup.json').read_text());assert len(backup['projects'])==3;assert backup['assessments']['A.5.1']['compliance']==5;assert backup['assessments']['A.5.1']['previousCompliance']=='5';assert backup['assessments']['4.1']['compliance']==5
 page.get_by_role('button',name='گزارش Word و Excel',exact=True).click();page.get_by_label('نام سازمان',exact=True).fill('سازمان آزمون نسخه۱۲');page.get_by_label('نوع ممیزی',exact=True).select_option(label='ممیزی داخلی');page.get_by_label('نام ممیز / بررسی‌کننده',exact=True).fill('ممیز آزمون')
 with page.expect_download() as d:page.get_by_role('button',name='دریافت Excel (.xlsx)',exact=True).click()
 d.value.save_as(str(art/'audit.xlsx'));wb=load_workbook(art/'audit.xlsx');sheet=wb['خلاصه یافته‌ها'];headers=[c.value for c in sheet[1]];good=headers.index('نقطه قوت (Good Point)')+1;assert int(str(sheet.cell(sheet.max_row,good).value).translate(str.maketrans('۰۱۲۳۴۵۶۷۸۹','0123456789')))==2
 project_sheet=wb['اقدامات بهبود'];project_headers=[c.value for c in project_sheet[1]];project_row=list(project_sheet.values)[1];num=lambda v:int(str(v).translate(str.maketrans('۰۱۲۳۴۵۶۷۸۹','0123456789')).replace(',','').replace('٬',''))
 for label,expected in [('بودجه مصوب (ریال)',1000),('جمع تخصیص (ریال)',350),('مانده بودجه (ریال)',650)]:assert num(project_row[project_headers.index(label)])==expected
 assert project_row[project_headers.index('گروه زمانی')]=='A';assert 'شناخت سازمان' in project_row[project_headers.index('سرفصل‌های مرتبط')];assert wb['تخصیص بودجه اقدامات اصلاحی'].max_row==3
 with page.expect_download() as d:page.get_by_role('button',name='دریافت Word (.docx)',exact=True).click()
 d.value.save_as(str(art/'audit.docx'))
 with ZipFile(art/'audit.docx') as z:
  assert z.testzip() is None;xml=z.read('word/document.xml').decode();assert 'Good Point' in xml
 page.get_by_label('بستن گزارش',exact=True).click();main('پشتیبانی سایت');page.get_by_label('بارگذاری نسخه JSON',exact=True).set_input_files(str(out/'backup.json'));page.get_by_label('نام سازمان واردشده',exact=True).fill('بازیابی نسخه۱۲');page.get_by_role('button',name='واردکردن پروژه',exact=True).click();expect(page.get_by_role('dialog')).to_have_count(0);page.reload();main('اقدامات اصلاحی');assert page.locator('[data-improvement-project]').count()==3
 page.get_by_role('button',name='ویرایش پروژه پروژه حاکمیت',exact=True).click();form=page.get_by_role('dialog');form.locator('summary').filter(has_text='بودجه اقدام اصلاحی').click();assert '۳۵۰' in form.locator('[data-v12-budget-total]').inner_text();assert '۶۵۰' in form.locator('[data-v12-budget-remaining]').inner_text();form.get_by_role('button',name='ذخیره پروژه',exact=True).click()
 # Styling and daily actions remain usable when disconnected, including the narrower viewport.
 ctx.set_offline(True);page.get_by_label('پوسته سایت',exact=True).select_option('dark');page.screenshot(path=str(art/'projects-dark.png'),full_page=True);page.set_viewport_size({'width':390,'height':844});assert not page.evaluate('document.documentElement.scrollWidth>innerWidth');page.get_by_role('button',name='ویرایش پروژه پروژه حاکمیت',exact=True).click();assert not page.evaluate('document.documentElement.scrollWidth>innerWidth');page.get_by_role('dialog').screenshot(path=str(art/'project-mobile.png'));page.get_by_role('button',name='ذخیره پروژه',exact=True).click();main('داشبورد مدیریتی');assert page.locator('[data-standard-dot="A.5.1"]').get_attribute('data-compliance')=='5'
 assert not errors,errors;assert all(r.split('#')[0]==args.url for r in requests),requests
 (art/'result.json').write_text(json.dumps({'result':'PASS','errors':errors,'requests':requests,'groupColors':colors},indent=2));print('PASS JSON backup/restore, Good Point Excel/Word reports, persisted projects and budgets, dark/mobile/offline with no JS errors or external requests.');b.close()
