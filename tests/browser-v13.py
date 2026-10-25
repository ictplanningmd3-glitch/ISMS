"""Offline v13 QA: shared taxonomy, ordered compliance, normalized bars and corrective dashboard."""
from playwright.sync_api import sync_playwright, expect
from openpyxl import load_workbook
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile
import argparse, json, math, os

p=argparse.ArgumentParser()
p.add_argument('--url',default='http://127.0.0.1:8773/Didban_Amniat_Offline.html')
p.add_argument('--artifacts',default='/tmp/didban-v13-qa')
args=p.parse_args()
art=Path(args.artifacts);art.mkdir(parents=True,exist_ok=True)
taxonomy=json.loads(Path('src/taxonomy-v13.json').read_text())
groups=taxonomy['groups'];clauses=taxonomy['clauses']
assert len(clauses)==10 and sum(len(c['itemIds']) for c in clauses)==30
assert len({id for c in clauses for id in c['itemIds']})==30
assert [g['id'] for g in groups]==['governance','assets','identity','people','physical','infrastructure','threat','data','development','supplier']
assert [len(g['controls']) for g in groups]==[15,6,9,8,14,5,8,6,10,12]
for control,domain in [('A.5.7','governance'),('A.5.12','assets'),('A.5.24','supplier'),('A.8.1','infrastructure'),('A.8.13','data'),('A.8.34','development')]:assert control in next(g for g in groups if g['id']==domain)['controls']
ids=[id for g in groups for id in g['controls']]
expected_ids=[f'A.{domain}.{n}' for domain,count in [(5,37),(6,8),(7,14),(8,34)] for n in range(1,count+1)]
assert len(ids)==93 and len(set(ids))==93 and set(ids)==set(expected_ids)
order=['5','0','4','3','1','2']
num=lambda v:int(str(v).translate(str.maketrans('۰۱۲۳۴۵۶۷۸۹','0123456789')).replace(',','').replace('٬',''))

with TemporaryDirectory() as temp,sync_playwright() as pw:
 out=Path(temp)
 b=pw.chromium.launch(executable_path=os.environ.get('DIDBAN_CHROMIUM','/usr/bin/chromium'),headless=True,args=['--no-sandbox','--disable-dev-shm-usage'])
 ctx=b.new_context(viewport={'width':1536,'height':1050},accept_downloads=True)
 page=ctx.new_page();errors=[];requests=[]
 page.on('pageerror',lambda e:errors.append(str(e)))
 page.on('request',lambda r:requests.append(r.url))
 page.goto(args.url)
 def main(name):page.locator('aside nav>button').filter(has_text=name).click()
 def corrective_tab(value):page.locator('[data-v13-corrective-tab="'+value+'"]:visible').first.click()
 def save_audit():
  page.get_by_role('button',name='ذخیره ارزیابی',exact=True).click()
  expect(page.get_by_role('dialog')).to_have_count(0)
 def capture(locator,name):
  style=page.add_style_tag(content='body .shell>main>header,body .toast{visibility:hidden!important}')
  locator.screenshot(path=str(art/name))
  style.evaluate('e=>e.remove()')
 def main_structure():
  assert page.locator('[data-v12-project-dashboard=main]').count()==0
  assert page.locator('.critical-panel,.recent,.bottomgrid').count()==0
  for text in ['آخرین ارزیابی‌ها','عملکرد حوزه‌های Annex A','آخرین وضعیت و پیشرفت پروژه‌های بهبود']:
   assert page.get_by_text(text,exact=True).count()==0,text
  assert page.locator('[data-standard-dot]').count()==93
  assert page.locator('[data-compliance-group]').count()==10
  assert page.locator('[data-v13-implementation-axis]').evaluate_all('xs=>xs.map(x=>x.dataset.v13ImplementationAxis)')==['clause:'+c['id'] for c in clauses]
  assert page.locator('[data-v13-compliance-group]').count()==20
  assert page.locator('.dc-section thead [data-compliance-column]').evaluate_all('xs=>xs.map(x=>x.dataset.complianceColumn)')==order+['6']
  for a in page.locator('.ds-quick a').all():
   href=a.get_attribute('href');assert href and href.startswith('#')
   assert page.locator('[id="'+href[1:]+'"]').count()==1,href
   a.click()
  assert len(page.locator('.ds-quick a').all())==6
  for empty_track in page.locator('.v13-nc-track.is-empty').all():assert empty_track.bounding_box()['height']<=45

 main_structure()
 assert page.locator('[data-v13-nc-row]').count()==10
 assert all(v=='0' for v in page.locator('[data-v13-nc-row]').evaluate_all('xs=>xs.map(x=>x.dataset.total)'))
 assert page.locator('[data-v13-nc-status]').count()==0
 # Good Point uses the same native menu appearance as the other choices.
 page.locator('[data-standard-dot="A.5.1"]').click()
 field=page.get_by_label('انطباق با استاندارد',exact=True)
 assert field.locator('option').evaluate_all('xs=>xs.map(x=>x.value)')==['']+order
 field.select_option('0');plain=field.evaluate('e=>getComputedStyle(e).backgroundColor')
 field.select_option('5');assert field.evaluate('e=>getComputedStyle(e).backgroundColor')==plain
 assert not field.get_attribute('class') or 'dc-good-point' not in field.get_attribute('class')
 good_option=field.locator('option[value="5"]')
 assert not good_option.get_attribute('style') and not good_option.get_attribute('class')
 page.get_by_label('نام ممیز',exact=True).fill('ممیز نسخه۱۳')
 page.get_by_text('اطلاعات ممیزی قبلی (در صورت وجود)',exact=True).click()
 prior=page.get_by_label('وضعیت انطباق در ممیزی قبلی',exact=True)
 prior.select_option('0');prior_plain=prior.evaluate('e=>getComputedStyle(e).backgroundColor')
 prior.select_option('5');assert prior.evaluate('e=>getComputedStyle(e).backgroundColor')==prior_plain
 save_audit();page.reload()
 assert page.locator('[data-standard-dot="A.5.1"]').get_attribute('data-compliance-rank')=='0'
 # Controlled records make normalization and gap calculations independently predictable.
 records=[('A.5.1',5,3,'2027-12-01',True),('A.5.2',0,3,'',True),('A.5.3',4,3,'',True),('A.5.4',3,2,'2027-12-01',True),('A.5.5',1,1,'2026-01-01',True),('A.5.6',2,0,'2026-01-01',True),('A.5.8',1,3,'2026-01-01',True),('A.5.31',0,3,'2027-12-01',True),('A.5.32',2,0,'2026-01-01',False),('A.6.1',1,1,'2027-12-01',True),('4.1',2,0,'2026-01-01',True),('4.2',1,2,'2026-01-01',True),('5.1',5,4,'2026-01-01',True)]
 page.evaluate("""rows=>{const reg=JSON.parse(localStorage.getItem('didban-security-workspaces-v1:default')),key='didban-project:default:'+reg.activeId+':didban-isms-offline-v1';const previous=JSON.parse(localStorage.getItem(key))['A.5.1'];localStorage.setItem(key,JSON.stringify(Object.fromEntries(rows.map(([id,compliance,status,due,applicable])=>[id,{...previous,id,compliance,status,due,applicable,owner:'مالک ممیزی',reason:applicable?'':'خارج از دامنه آزمون',action:'اقدام ثبت‌شده',updated:'2026-10-25T08:00:00Z'}]))));}""",records)
 page.reload();main_structure()
 assert page.locator('[data-v13-implementation-score="clause:interested-scope"]').get_attribute('data-score')=='17'
 assert page.locator('[data-v13-implementation-score="clause:leadership"]').get_attribute('data-score')=='33'
 # Rank and the actual radius both follow the requested center-to-edge order.
 colors=[];radii=[]
 for rank,(id,value) in enumerate([('A.5.1',5),('A.5.2',0),('A.5.3',4),('A.5.4',3),('A.5.5',1),('A.5.6',2)]):
  dot=page.locator('[data-standard-dot="'+id+'"]')
  assert dot.get_attribute('data-compliance')==str(value)
  assert dot.get_attribute('data-compliance-rank')==str(rank)
  x,y,color=dot.evaluate('e=>[+e.getAttribute("cx"),+e.getAttribute("cy"),getComputedStyle(e).fill]')
  radii.append(math.hypot(x-260,y-260));colors.append(color)
 assert radii==sorted(radii) and max(radii)<190 and len(set(colors))==6
 assert page.locator('.dc-summary article').filter(has_text='Good Point').locator('strong').inner_text()=='۲'
 for i,g in enumerate(groups):
  card=page.locator('[data-compliance-group="'+g['id']+'"]')
  assert g['title'] in card.inner_text() and g['english'] in card.inner_text()
  assert card.count()==1
 # Two actual findings in governance and one major normalize to 2/3 and 1/3, not count width.
 governance=page.locator('[data-v13-nc-row="annex:governance"]')
 assert [governance.get_attribute('data-'+k) for k in ['minor','major','total']]==['2','1','3']
 widths=governance.locator('[data-v13-nc-status]').evaluate_all('xs=>xs.map(x=>parseFloat(x.style.width))')
 assert len(widths)==2 and abs(sum(widths)-100)<0.01 and abs(widths[0]-200/3)<0.01 and abs(widths[1]-100/3)<0.01,widths
 assert governance.locator('[data-v13-nc-status]').evaluate_all('xs=>xs.map(x=>x.dataset.v13NcStatus)')==['1','2']
 people=page.locator('[data-v13-nc-row="annex:people"]')
 assert people.get_attribute('data-total')=='1'
 assert people.locator('[data-v13-nc-status="1"]').evaluate('e=>parseFloat(e.style.width)')==100
 empty=page.locator('[data-v13-nc-row="annex:assets"]');assert empty.get_attribute('data-total')=='0' and empty.locator('[data-v13-nc-status]').count()==0
 governance.locator('[data-v13-nc-status]').first.hover()
 capture(page.locator('[data-v13-nc-dashboard]'),'nonconformity-annex-light.png')
 page.get_by_label('دامنه توزیع عدم‌انطباق‌ها',exact=True).select_option('clause')
 assert page.locator('[data-v13-nc-row]').count()==10
 context=page.locator('[data-v13-nc-row="clause:context"]')
 assert [context.get_attribute('data-'+k) for k in ['minor','major','total']]==['0','1','1']
 assert context.locator('[data-v13-nc-status]').evaluate_all('xs=>xs.map(x=>parseFloat(x.style.width))')==[100]
 interested=page.locator('[data-v13-nc-row="clause:interested-scope"]');assert [interested.get_attribute('data-'+k) for k in ['minor','major','total']]==['1','0','1']
 capture(page.locator('[data-v13-nc-dashboard]'),'nonconformity-clauses-light.png')
 print('PASS plain current/prior menus, six ranks/colors/order, shared ten-domain radar, removed main panels/bookmarks, real normalized Annex/clause NC bars',flush=True)
 # Group-card navigation applies the same canonical mapping to the control register.
 for g in groups:
  main('داشبورد مدیریتی');page.locator('[data-compliance-group="'+g['id']+'"]').click()
  assert page.get_by_label('حوزه مفهومی',exact=True).locator('option').count()==11
  table_ids=page.locator('.tablewrap tbody tr').evaluate_all('rows=>rows.map(r=>r.querySelector("bdi")?.textContent).filter(Boolean)')
  assert set(table_ids)==set(g['controls']),(g['id'],table_ids,g['controls'])
 for c in clauses:
  main('الزامات استاندارد');selector=page.get_by_label('حوزه مفهومی',exact=True)
  assert selector.locator('option').count()==11
  selector.select_option(c['id'])
  table_ids=page.locator('.tablewrap tbody tr').evaluate_all('rows=>rows.map(r=>r.querySelector("bdi")?.textContent).filter(Boolean)')
  assert set(table_ids)==set(c['itemIds']),(c['id'],table_ids,c['itemIds'])
 for heading,id in [(groups[0]['english'],'A.5.1'),(clauses[0]['english'],'4.1')]:
  page.get_by_label('جستجو در محتوای سایت',exact=True).fill(heading)
  page.locator('.ds-search-results>button').filter(has_text=id+' —').first.click()
  expect(page.get_by_role('dialog')).to_have_count(1)
  page.get_by_role('dialog').get_by_role('button',name='بستن',exact=True).click()
 # Corrective portfolio still retains the useful project and budget workflow.
 main('اقدامات اصلاحی');assert page.locator('aside [data-v13-corrective-tab]').count()==2
 assert page.locator('[data-v13-scope-bars="clause"]').count()==1 and page.locator('[data-v12-scope-radar]').count()==0
 projects=[('پروژه حاکمیت','1','active','25',[clauses[0]['title'],clauses[1]['title'],groups[0]['title']],'۱۴۰۴/۱۰/۱۱'),('پروژه افراد','2','blocked','60',[groups[3]['title']],'۱۴۰۵/۰۸/۰۲'),('پروژه رهبری','3','completed','100',[clauses[2]['title']],'۱۴۰۴/۱۰/۱۱')]
 for i,(title,months,status,progress,scopes,start) in enumerate(projects):
  page.get_by_role('button',name='تعریف اقدام بهبود',exact=True).click();form=page.get_by_role('dialog')
  for label,value in [('عنوان اقدام یا پروژه',title),('مسئول پیگیری','مالک '+title),('هدف و نتیجه مورد انتظار','رفع شکاف استاندارد'),('منابع لازم','تیم اجرا')]:form.get_by_label(label,exact=True).fill(value)
  for g in groups+clauses:assert form.get_by_label('حوزه پروژه '+g['title'],exact=True).count()==1
  form.get_by_label('گروه زمانی پروژه',exact=True).select_option(months)
  form.get_by_label('وضعیت پروژه',exact=True).select_option(status)
  if status!='completed':form.get_by_label('پیشرفت (درصد)',exact=True).fill(progress)
  form.get_by_label('تاریخ شروع پروژه',exact=True).fill(start)
  for heading in scopes:form.get_by_label('حوزه پروژه '+heading,exact=True).check()
  if i==1:
   form.locator('summary').filter(has_text='ارتباط دقیق با الزامات و کنترل‌ها').click();form.get_by_label('ارتباط پروژه A.5.5',exact=True).check()
  if i==0:
   form.locator('summary').filter(has_text='بودجه اقدام اصلاحی').click()
   form.get_by_label('بودجه مصوب اقدام اصلاحی',exact=True).fill('1000')
   form.get_by_label('تاریخ تصویب بودجه',exact=True).fill('۱۴۰۵/۰۸/۰۱')
   for j,(amount,date) in enumerate([('100','۱۴۰۵/۰۸/۰۲'),('250','۱۴۰۵/۰۸/۰۳')],1):
    form.get_by_role('button',name='افزودن ردیف تخصیص بودجه',exact=True).click()
    form.get_by_label('میزان بودجه تخصیص داده شده '+str(j),exact=True).fill(amount)
    form.get_by_label('تاریخ تخصیص بودجه '+str(j),exact=True).fill(date)
   assert '۳۵۰' in form.locator('[data-v12-budget-total]').inner_text() and '۶۵۰' in form.locator('[data-v12-budget-remaining]').inner_text()
  form.get_by_role('button',name='ذخیره پروژه',exact=True).click();expect(form).to_have_count(0)
 assert page.locator('[data-improvement-project]').count()==3
 assert page.locator('[data-v12-scope="clause:context"]').get_attribute('data-count')=='1'
 page.get_by_label('دامنه نمودار گستره پروژه‌ها',exact=True).select_option('control')
 assert page.locator('[data-v13-scope-bars="control"]').count()==1
 assert page.locator('[data-v12-scope="annex:governance"]').get_attribute('data-count')=='2'
 assert page.locator('[data-v12-scope="annex:people"]').get_attribute('data-count')=='1'
 assert '۶۲٪' in page.locator('.v12-average').inner_text()
 colors=[page.locator('[data-project-group="'+g+'"][data-improvement-project]').evaluate('e=>getComputedStyle(e).borderTopColor') for g in ['A','B','C']]
 assert len(set(colors))==3
 capture(page.locator('.v12-project-charts'),'project-bars-light.png')
 corrective_tab('gaps')
 assert page.locator('[data-v13-corrective-dashboard]').count()==1
 assert page.locator('[data-improvement-project]').count()==0
 for key,value in [('critical',5),('overdue-audits',4),('overdue-projects',1),('uncovered',0),('major',2),('minor',4)]:
  assert num(page.locator('[data-v13-kpi="'+key+'"]').locator('strong').inner_text())==value,key
 assert page.locator('[data-v13-audit-attention]').count()==6
 assert set(page.locator('[data-v13-gap-chart]').evaluate_all('xs=>xs.map(x=>x.dataset.v13GapChart)'))=={'scope','coverage','deadlines','projects'}
 assert page.locator('[data-v13-gap-chart="coverage"] svg').count()==1
 page.get_by_label('فیلتر شکاف‌ها و اقدامات معوق',exact=True).select_option('critical');assert page.locator('[data-v13-audit-attention]').count()==5
 page.get_by_label('فیلتر شکاف‌ها و اقدامات معوق',exact=True).select_option('overdue');assert page.locator('[data-v13-audit-attention]').count()==4
 page.get_by_label('فیلتر شکاف‌ها و اقدامات معوق',exact=True).select_option('uncovered');assert page.locator('[data-v13-audit-attention]').count()==0
 page.get_by_label('فیلتر شکاف‌ها و اقدامات معوق',exact=True).select_option('all')
 page.get_by_role('button',name='بازکردن ارزیابی A.5.6',exact=True).click()
 expect(page.get_by_role('dialog')).to_have_count(1)
 assert page.get_by_label('انطباق با استاندارد',exact=True).input_value()=='2'
 page.get_by_role('dialog').get_by_role('button',name='بستن',exact=True).click()
 page.get_by_role('button',name='بازکردن پروژه مرتبط پروژه حاکمیت',exact=True).first.click()
 expect(page.get_by_label('عنوان اقدام یا پروژه',exact=True)).to_have_value('پروژه حاکمیت')
 page.get_by_role('button',name='انصراف',exact=True).click()
 page.get_by_role('button',name='تعریف اقدام برای A.5.8',exact=True).click()
 assert page.get_by_label('ارتباط پروژه A.5.8',exact=True).is_checked()
 page.get_by_role('button',name='انصراف',exact=True).click()
 capture(page.locator('[data-v13-corrective-dashboard]'),'corrective-gaps-light.png')
 print('PASS canonical controls navigation and project form, scope bars, A/B/C cards, computed project progress and budget',flush=True)
 # Native portable backup and Word/Excel both preserve the changed data and status order.
 main('پشتیبانی سایت')
 with page.expect_download() as d:page.get_by_role('button',name='دریافت JSON پروژه فعال',exact=True).click()
 d.value.save_as(str(out/'backup.json'));backup=json.loads((out/'backup.json').read_text())
 assert len(backup['projects'])==3 and backup['assessments']['A.5.1']['compliance']==5 and backup['assessments']['A.5.1']['previousCompliance']=='5'
 page.get_by_role('button',name='گزارش Word و Excel',exact=True).click()
 page.get_by_label('نام سازمان',exact=True).fill('سازمان آزمون نسخه۱۳')
 page.get_by_label('نوع ممیزی',exact=True).select_option(label='ممیزی داخلی')
 page.get_by_label('نام ممیز / بررسی‌کننده',exact=True).fill('ممیز آزمون')
 with page.expect_download() as d:page.get_by_role('button',name='دریافت Excel (.xlsx)',exact=True).click()
 d.value.save_as(str(art/'audit.xlsx'));wb=load_workbook(art/'audit.xlsx')
 sheet=wb['خلاصه یافته‌ها'];headers=[c.value for c in sheet[1]]
 assert headers[1:7]==['نقطه قوت (Good Point)','انطباق','توصیه برای بهبود','مشاهده','عدم انطباق جزیی','عدم انطباق کلی']
 assert num(sheet.cell(sheet.max_row,headers.index('نقطه قوت (Good Point)')+1).value)==2
 project_sheet=wb['اقدامات بهبود'];project_headers=[c.value for c in project_sheet[1]];row=list(project_sheet.values)[1]
 for name,value in [('بودجه مصوب (ریال)',1000),('جمع تخصیص (ریال)',350),('مانده بودجه (ریال)',650)]:assert num(row[project_headers.index(name)])==value
 assert row[project_headers.index('گروه زمانی')]=='A' and wb['تخصیص بودجه اقدامات اصلاحی'].max_row==3
 with page.expect_download() as d:page.get_by_role('button',name='دریافت Word (.docx)',exact=True).click()
 d.value.save_as(str(art/'audit.docx'))
 with ZipFile(art/'audit.docx') as z:assert z.testzip() is None and 'Good Point' in z.read('word/document.xml').decode()
 page.get_by_label('بستن گزارش',exact=True).click();main('پشتیبانی سایت')
 page.get_by_label('بارگذاری نسخه JSON',exact=True).set_input_files(str(out/'backup.json'))
 page.get_by_label('نام سازمان واردشده',exact=True).fill('بازیابی نسخه۱۳')
 page.get_by_role('button',name='واردکردن پروژه',exact=True).click();expect(page.get_by_role('dialog')).to_have_count(0)
 page.reload();main('اقدامات اصلاحی');corrective_tab('projects')
 assert page.locator('[data-improvement-project]').count()==3
 page.get_by_role('button',name='ویرایش پروژه پروژه حاکمیت',exact=True).click();form=page.get_by_role('dialog')
 form.locator('summary').filter(has_text='بودجه اقدام اصلاحی').click()
 assert '۳۵۰' in form.locator('[data-v12-budget-total]').inner_text() and '۶۵۰' in form.locator('[data-v12-budget-remaining]').inner_text()
 form.get_by_role('button',name='ذخیره پروژه',exact=True).click()
 # The full offline runtime, both themes, and narrow layouts remain usable.
 ctx.set_offline(True);page.get_by_label('پوسته سایت',exact=True).select_option('dark')
 corrective_tab('gaps');capture(page.locator('[data-v13-corrective-dashboard]'),'corrective-gaps-dark.png')
 page.set_viewport_size({'width':390,'height':844})
 assert not page.evaluate('document.documentElement.scrollWidth>innerWidth')
 capture(page.locator('[data-v13-corrective-dashboard]'),'corrective-gaps-mobile.png')
 main('داشبورد مدیریتی');main_structure()
 assert not page.evaluate('document.documentElement.scrollWidth>innerWidth')
 assert page.locator('[data-v13-nc-status]>span').first.evaluate('e=>getComputedStyle(e).color')=='rgb(11, 18, 32)'
 capture(page.locator('[data-v13-nc-dashboard]'),'nonconformity-mobile.png')
 main('اقدامات اصلاحی');corrective_tab('projects')
 page.get_by_role('button',name='ویرایش پروژه پروژه حاکمیت',exact=True).click()
 assert not page.evaluate('document.documentElement.scrollWidth>innerWidth')
 capture(page.get_by_role('dialog'),'project-mobile.png')
 page.get_by_role('button',name='ذخیره پروژه',exact=True).click()
 assert not errors,errors
 assert all(r.split('#')[0]==args.url for r in requests),requests
 (art/'result.json').write_text(json.dumps({'result':'PASS','errors':errors,'requests':requests,'taxonomyCount':len(groups),'clauseGroups':len(clauses),'controls':len(ids),'radii':radii},indent=2))
 print('PASS project backup/restore, status order in Excel/Word, budget and group exports, dark/mobile/offline, no JS errors or external requests.');b.close()
