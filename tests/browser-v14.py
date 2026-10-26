"""Offline v14 QA: all audit results, eight domains, and automatic corrective states."""
from playwright.sync_api import sync_playwright, expect
from openpyxl import load_workbook
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile
import argparse, json, math, os

p=argparse.ArgumentParser()
p.add_argument('--url',default='http://127.0.0.1:8774/Didban_Amniat_Offline.html')
p.add_argument('--artifacts',default='/tmp/didban-v14-qa')
args=p.parse_args();art=Path(args.artifacts);art.mkdir(parents=True,exist_ok=True)
taxonomy=json.loads(Path('src/taxonomy-v13.json').read_text())
groups=taxonomy['groups'];clauses=taxonomy['clauses'];order=['5','0','4','3','1','2'];findings=['4','3','1','2']
assert len(groups)==10 and len(clauses)==8
assert clauses[0]['title']=='شناخت سازمان، دامنه و ذینفعان' and clauses[0]['itemIds']==['4.1','4.2','4.3','4.4']
assert clauses[2]['title']=='مدیریت ریسک' and clauses[2]['itemIds']==['6.1.1','6.1.2','6.1.3']
assert len({x for g in groups for x in g['controls']})==93
assert len({x for g in clauses for x in g['itemIds']})==30
num=lambda value:int(str(value).translate(str.maketrans('۰۱۲۳۴۵۶۷۸۹','0123456789')).replace('٬','').replace(',',''))

with TemporaryDirectory() as temp,sync_playwright() as pw:
 out=Path(temp)
 browser=pw.chromium.launch(executable_path=os.environ.get('DIDBAN_CHROMIUM','/usr/bin/chromium'),headless=True,args=['--no-sandbox','--disable-dev-shm-usage'])
 ctx=browser.new_context(viewport={'width':1536,'height':1050},accept_downloads=True)
 page=ctx.new_page();errors=[];requests=[]
 page.on('pageerror',lambda e:errors.append(str(e)))
 page.on('request',lambda r:requests.append(r.url))
 page.goto(args.url)
 def main(name):page.locator('aside nav>button').filter(has_text=name).click()
 def tab(value):page.locator('[data-v13-corrective-tab="'+value+'"]:visible').first.click()
 def dashboard():main('اقدامات اصلاحی');tab('gaps')
 def row(item):return page.locator('[data-v14-action-row="'+item+'"]')
 def raw(kind):return page.evaluate("suffix=>{const key=Object.keys(localStorage).find(k=>k.endsWith(':'+suffix));return key?JSON.parse(localStorage.getItem(key)):null;}",kind)
 def capture(locator,name):
  style=page.add_style_tag(content='body .shell>main>header,body .toast{visibility:hidden!important}')
  locator.screenshot(path=str(art/name));style.evaluate('e=>e.remove()')
 def kpis(expected):
  assert page.locator('[data-v14-kpi]').count()==4
  for key,value in expected.items():assert int(page.locator('[data-v14-kpi="'+key+'"]').get_attribute('data-count'))==value,(key,value)
 def project_counts(overdue,scheduled,completed,clause,control):
  for key,value in [('overdue',overdue),('scheduled',scheduled),('completed',completed)]:
   assert int(page.locator('[data-v13-deadline-bucket="'+key+'"]').get_attribute('data-count'))==value,key
  assert page.locator('[data-v13-deadline-bucket]').count()==3
  for key,value in [('clause',clause),('control',control)]:
   assert int(page.locator('[data-v14-completed-kind="'+key+'"]').get_attribute('data-count'))==value,key
 def structure():
  assert page.locator('[data-standard-dot]').count()==93
  assert page.locator('[data-v13-implementation-axis]').count()==8
  assert page.locator('[data-v13-compliance-group]').count()==18
  assert page.locator('.dc-section thead [data-compliance-column]').evaluate_all('xs=>xs.map(x=>x.dataset.complianceColumn)')==order+['6']
  for link in page.locator('.ds-quick a').all():
   href=link.get_attribute('href');assert page.locator('[id="'+href[1:]+'"]').count()==1,href
  expect(page.get_by_role('heading',name='توزیع نتایج ممیزی',exact=True)).to_be_visible()
  assert page.locator('[data-v14-audit-legend]').evaluate_all('xs=>xs.map(x=>x.dataset.v14AuditLegend)')==order
  assert page.locator('.critical-panel,.recent,.bottomgrid,[data-v12-project-dashboard=main]').count()==0

 structure()
 assert page.locator('[data-v14-audit-series]').count()==0
 for empty in page.locator('.v13-nc-track.is-empty').all():assert empty.bounding_box()['height']<=45
 # Results requiring action must be determined by compliance, even with implementation status 4.
 audits=[('A.5.1',5,4,True),('A.5.2',0,0,True),('A.5.3',4,4,True),('A.5.4',3,4,True),('A.5.5',1,4,True),('A.5.6',2,4,True),('A.5.7',None,0,True),('A.5.8',2,0,False),('A.6.1',4,4,True),('4.1',3,4,True),('4.2',1,4,True),('6.1.1',2,4,True),('6.1.2',0,0,True),('5.1',5,4,True)]
 page.evaluate("""rows=>{const reg=JSON.parse(localStorage.getItem('didban-security-workspaces-v1:default')),prefix='didban-project:default:'+reg.activeId+':';localStorage.setItem(prefix+'didban-isms-offline-v1',JSON.stringify(Object.fromEntries(rows.map(([id,compliance,status,applicable])=>[id,{id,compliance,status,applicable,reason:applicable?'':'خارج از دامنه آزمون',documents:'شواهد بررسی‌شده',observations:'یافته آزمون',owner:'مالک آزمون',action:'تصمیم ثبت‌شده',due:'2026-01-01',reviewer:'ممیز آزمون',correctiveAction:'',previousCompliance:'',updated:'2026-10-26T08:00:00Z'}]))));localStorage.setItem(prefix+'didban-isms-improvements-v1','[]');}""",audits)
 page.reload();structure()
 gov=page.locator('[data-v14-audit-row="annex:governance"]')
 assert gov.get_attribute('data-total')=='6' and gov.get_attribute('data-unknown')=='8'
 assert gov.locator('[data-v14-audit-series]').evaluate_all('xs=>xs.map(x=>x.dataset.v14AuditSeries)')==order
 assert gov.locator('[data-v14-audit-series]').evaluate_all('xs=>xs.map(x=>x.dataset.count)')==['1']*6
 percentages=gov.locator('[data-v14-audit-series]').evaluate_all('xs=>xs.map(x=>parseFloat(x.style.width))')
 assert abs(sum(percentages)-100)<.01 and all(abs(v-100/6)<.01 for v in percentages),percentages
 radar_yellow=page.locator('[data-standard-dot="A.5.3"]').evaluate('e=>getComputedStyle(e).fill')
 recommendation=gov.locator('[data-v14-audit-series="4"]')
 assert recommendation.evaluate('e=>getComputedStyle(e).backgroundColor')==radar_yellow
 assert recommendation.locator(':scope>span').evaluate('e=>getComputedStyle(e).color')=='rgb(11, 18, 32)'
 capture(page.locator('[data-v14-audit-dashboard]'),'audit-results-annex-light.png')
 page.get_by_label('دامنه توزیع نتایج ممیزی',exact=True).select_option('clause')
 assert page.locator('[data-v14-audit-row]').count()==8
 context=page.locator('[data-v14-audit-row="clause:context"]')
 assert context.get_attribute('data-total')=='2'
 assert context.locator('[data-v14-audit-series]').evaluate_all('xs=>xs.map(x=>x.dataset.v14AuditSeries)')==['3','1']
 assert context.locator('[data-v14-audit-series]').evaluate_all('xs=>xs.map(x=>parseFloat(x.style.width))')==[50,50]
 capture(page.locator('[data-v14-audit-dashboard]'),'audit-results-clauses-light.png')
 # Shared catalog domains are identical in each live filter.
 for group in groups:
  main('داشبورد مدیریتی');page.locator('[data-compliance-group="'+group['id']+'"]').click()
  ids=page.locator('.tablewrap tbody tr').evaluate_all('xs=>xs.map(x=>x.querySelector("bdi")?.textContent).filter(Boolean)')
  assert set(ids)==set(group['controls']),group['id']
 for group in clauses:
  main('الزامات استاندارد');field=page.get_by_label('حوزه مفهومی',exact=True)
  assert field.locator('option').count()==9
  field.select_option(group['id'])
  ids=page.locator('.tablewrap tbody tr').evaluate_all('xs=>xs.map(x=>x.querySelector("bdi")?.textContent).filter(Boolean)')
  assert set(ids)==set(group['itemIds']),group['id']
 print('PASS all six audit results in radar order, real normalized percentages, unknown/nonapplicable exclusion, eight-domain merge and all 18 filters',flush=True)
 dashboard();dash=page.locator('[data-v13-corrective-dashboard]')
 expect(page.locator('[data-v13-corrective-tab=gaps]:visible').first).to_have_text('داشبورد اقدامات اصلاحی')
 kpis({'attention':8,'projects':0,'completed':0,'incomplete':0})
 assert page.locator('.v14-action-register thead th').all_text_contents()==['الزام / کنترل','وضعیت ممیزی','پروژه‌های بهبود تعریف‌شده','وضعیت اقدام']
 attention={'A.5.3','A.5.4','A.5.5','A.5.6','A.6.1','4.1','4.2','6.1.1'}
 assert set(page.locator('[data-v14-action-row]').evaluate_all('xs=>xs.map(x=>x.dataset.v14ActionRow)'))==attention
 assert all(s=='unplanned' for s in page.locator('[data-v14-action-row]').evaluate_all('xs=>xs.map(x=>x.dataset.v14ActionStatus)'))
 assert dash.get_by_role('button').filter(has_text='تعریف اقدام').count()==0
 assert dash.locator('[data-v14-action-row] td').count()==32
 assert page.get_by_text('پوشش پروژه‌های مرتبط',exact=True).count()==0
 assert page.locator('[data-v14-finding-chart=control] [data-v14-corrective-scope]').count()==10
 assert page.locator('[data-v14-finding-chart=clause] [data-v14-corrective-scope]').count()==8
 stack=page.locator('[data-v14-corrective-scope="annex:governance"]')
 assert stack.get_attribute('data-total')=='4'
 assert stack.locator('[data-v14-finding-status]').evaluate_all('xs=>xs.map(x=>x.dataset.v14FindingStatus)')==findings
 assert stack.locator('[data-v14-finding-status]').evaluate_all('xs=>xs.map(x=>parseFloat(x.style.width))')==[25]*4
 assert stack.locator('[data-v14-finding-status="4"]').evaluate('e=>getComputedStyle(e).backgroundColor')==radar_yellow
 assert stack.locator('[data-v14-finding-status="4"]>span').evaluate('e=>getComputedStyle(e).color')=='rgb(11, 18, 32)'
 project_counts(0,0,0,0,0)
 # Create projects through the portfolio, never through the corrective dashboard.
 definitions=[('پروژه مشترک فعال','active',['A.5.3','4.1'],[],'2026-01-01'),('پروژه بسته دوم','completed',['A.5.3'],[],'2026-01-01'),('پروژه حوزه‌ای','active',['A.6.1'],['annex:people','clause:risk-management'],'2027-01-01'),('پروژه تکمیل جزئی','completed',['A.5.5'],[],'2026-01-01'),('پروژه تکمیل مشترک','completed',['A.5.6','4.2'],[],'2026-01-01')]
 scope_titles={**{'annex:'+g['id']:g['title'] for g in groups},**{'clause:'+g['id']:g['title'] for g in clauses}}
 for i,(title,status,exact,scopes,start) in enumerate(definitions):
  tab('projects');page.get_by_role('button',name='تعریف اقدام بهبود',exact=True).click();form=page.get_by_role('dialog')
  for label,value in [('عنوان اقدام یا پروژه',title),('مسئول پیگیری','مالک '+title),('هدف و نتیجه مورد انتظار','رفع یافته ثبت‌شده'),('منابع لازم','تیم اجرا')]:form.get_by_label(label,exact=True).fill(value)
  assert form.locator('.v12-scope-options input[type=checkbox]').count()==18
  form.get_by_label('گروه زمانی پروژه',exact=True).select_option(str(i%3+1))
  form.get_by_label('وضعیت پروژه',exact=True).select_option(status)
  if status!='completed':form.get_by_label('پیشرفت (درصد)',exact=True).fill('25')
  form.get_by_label('تاریخ شروع پروژه',exact=True).fill(start)
  for scope in scopes:form.get_by_label('حوزه پروژه '+scope_titles[scope],exact=True).check()
  form.locator('summary').filter(has_text='ارتباط دقیق با الزامات و کنترل‌ها').click()
  for item in exact:form.get_by_label('ارتباط پروژه '+item,exact=True).check()
  if i==0:
   form.locator('summary').filter(has_text='بودجه اقدام اصلاحی').click()
   form.get_by_label('بودجه مصوب اقدام اصلاحی',exact=True).fill('1000')
   form.get_by_label('تاریخ تصویب بودجه',exact=True).fill('۱۴۰۵/۰۸/۰۱')
   for j,(amount,date) in enumerate([('100','۱۴۰۵/۰۸/۰۲'),('250','۱۴۰۵/۰۸/۰۳')],1):
    form.get_by_role('button',name='افزودن ردیف تخصیص بودجه',exact=True).click()
    form.get_by_label('میزان بودجه تخصیص داده شده '+str(j),exact=True).fill(amount)
    form.get_by_label('تاریخ تخصیص بودجه '+str(j),exact=True).fill(date)
   assert '۳۵۰' in form.locator('[data-v12-budget-total]').inner_text() and '۶۵۰' in form.locator('[data-v12-budget-remaining]').inner_text()
  form.get_by_role('button',name='ذخیره پروژه',exact=True).click();expect(form).to_have_count(0)
  if i==0:
   tab('gaps');assert row('A.5.3').get_attribute('data-v14-action-status')=='active'
 tab('gaps');kpis({'attention':8,'projects':5,'completed':3,'incomplete':2})
 project_counts(1,1,3,1,3)
 assert row('A.5.3').get_attribute('data-v14-action-status')=='active'
 assert row('A.5.3').get_by_role('button',name='بازکردن پروژه مرتبط پروژه بسته دوم',exact=True).count()==1
 assert row('A.5.4').get_attribute('data-v14-action-status')=='unplanned'
 for item in ['A.5.5','A.5.6','4.2']:assert row(item).get_attribute('data-v14-action-status')=='done'
 assert row('A.6.1').get_by_role('button',name='بازکردن پروژه مرتبط پروژه حوزه‌ای',exact=True).count()==1
 assert row('6.1.1').get_attribute('data-v14-action-status')=='active'
 for code,count in [('4',2),('3',2),('1',2),('2',2)]:
  page.get_by_label('فیلتر نتیجه ممیزی',exact=True).select_option(code);assert page.locator('[data-v14-action-row]').count()==count
 page.get_by_label('فیلتر نتیجه ممیزی',exact=True).select_option('all')
 page.get_by_label('فیلتر وضعیت اقدام',exact=True).select_option('unplanned');assert page.locator('[data-v14-action-row]').count()==1
 page.get_by_label('فیلتر وضعیت اقدام',exact=True).select_option('all')
 capture(dash,'corrective-dashboard-light.png')
 # Completion changes automatic action state only after every linked project is complete.
 row('A.5.3').get_by_role('button',name='بازکردن پروژه مرتبط پروژه مشترک فعال',exact=True).click();form=page.get_by_role('dialog')
 form.get_by_label('وضعیت پروژه',exact=True).select_option('completed')
 form.get_by_role('button',name='ذخیره پروژه',exact=True).click();expect(form).to_have_count(0)
 assert row('A.5.3').get_attribute('data-v14-action-status')=='done'
 assert row('4.1').get_attribute('data-v14-action-status')=='done'
 kpis({'attention':8,'projects':5,'completed':4,'incomplete':1});project_counts(0,1,4,2,4)
 assert raw('didban-isms-offline-v1')['A.5.3']['compliance']==4
 # Reopening one completed project reverses only affected automatic states.
 row('A.5.3').get_by_role('button',name='بازکردن پروژه مرتبط پروژه بسته دوم',exact=True).click();form=page.get_by_role('dialog')
 form.get_by_label('وضعیت پروژه',exact=True).select_option('active');form.get_by_label('پیشرفت (درصد)',exact=True).fill('50')
 form.get_by_role('button',name='ذخیره پروژه',exact=True).click();expect(form).to_have_count(0)
 assert row('A.5.3').get_attribute('data-v14-action-status')=='active'
 assert row('4.1').get_attribute('data-v14-action-status')=='done'
 kpis({'attention':8,'projects':5,'completed':3,'incomplete':2});project_counts(1,1,3,2,3)
 assert 'توصیه برای بهبود' in row('A.5.3').locator('td').nth(1).inner_text()
 assert row('A.5.3').locator('select').count()==0
 page.reload();dashboard();assert row('A.5.3').get_attribute('data-v14-action-status')=='active'
 assert row('4.1').get_attribute('data-v14-action-status')=='done'
 row('A.5.3').get_by_role('button',name='بازکردن ارزیابی A.5.3',exact=True).click()
 assert page.get_by_label('انطباق با استاندارد',exact=True).input_value()=='4'
 page.get_by_role('dialog').get_by_role('button',name='بستن',exact=True).click()
 print('PASS four independent finding types/KPIs, four-column table, no direct create, all related projects, automatic none/active/done and reopening, both fixed finding stacks/project deadlines/completed-kind counts',flush=True)
 # Portable backup and reports preserve all project links, completed states and budget data.
 main('پشتیبانی سایت')
 with page.expect_download() as download:page.get_by_role('button',name='دریافت JSON پروژه فعال',exact=True).click()
 download.value.save_as(str(out/'backup.json'));backup=json.loads((out/'backup.json').read_text())
 assert len(backup['projects'])==5 and backup['assessments']['A.5.3']['compliance']==4
 page.get_by_role('button',name='گزارش Word و Excel',exact=True).click()
 page.get_by_label('نام سازمان',exact=True).fill('سازمان آزمون نسخه۱۴')
 page.get_by_label('نوع ممیزی',exact=True).select_option(label='ممیزی داخلی')
 page.get_by_label('نام ممیز / بررسی‌کننده',exact=True).fill('ممیز آزمون')
 with page.expect_download() as download:page.get_by_role('button',name='دریافت Excel (.xlsx)',exact=True).click()
 download.value.save_as(str(art/'audit.xlsx'));workbook=load_workbook(art/'audit.xlsx')
 summary=workbook['خلاصه یافته‌ها'];assert summary.max_row==20
 assert [c.value for c in summary[1]][1:7]==['نقطه قوت (Good Point)','انطباق','توصیه برای بهبود','مشاهده','عدم انطباق جزیی','عدم انطباق کلی']
 assert all(num(summary.cell(summary.max_row,i).value)==2 for i in range(2,8))
 project_sheet=workbook['اقدامات بهبود'];headers=[c.value for c in project_sheet[1]];values=list(project_sheet.values)[1]
 for title,value in [('بودجه مصوب (ریال)',1000),('جمع تخصیص (ریال)',350),('مانده بودجه (ریال)',650)]:assert num(values[headers.index(title)])==value
 with page.expect_download() as download:page.get_by_role('button',name='دریافت Word (.docx)',exact=True).click()
 download.value.save_as(str(art/'audit.docx'))
 with ZipFile(art/'audit.docx') as archive:
  assert archive.testzip() is None;xml=archive.read('word/document.xml').decode()
  assert 'شناخت سازمان، دامنه و ذینفعان' in xml and 'مدیریت ریسک' in xml
 page.get_by_label('بستن گزارش',exact=True).click();main('پشتیبانی سایت')
 page.get_by_label('بارگذاری نسخه JSON',exact=True).set_input_files(str(out/'backup.json'))
 page.get_by_label('نام سازمان واردشده',exact=True).fill('بازیابی نسخه۱۴')
 page.get_by_role('button',name='واردکردن پروژه',exact=True).click();expect(page.get_by_role('dialog')).to_have_count(0)
 page.reload();dashboard();kpis({'attention':8,'projects':5,'completed':3,'incomplete':2})
 assert row('A.5.3').get_attribute('data-v14-action-status')=='active' and row('4.1').get_attribute('data-v14-action-status')=='done'
 ctx.set_offline(True);page.get_by_label('پوسته سایت',exact=True).select_option('dark')
 capture(page.locator('[data-v13-corrective-dashboard]'),'corrective-dashboard-dark.png')
 main('داشبورد مدیریتی');structure()
 page.get_by_label('دامنه توزیع نتایج ممیزی',exact=True).select_option('control')
 capture(page.locator('[data-v14-audit-dashboard]'),'audit-results-dark.png')
 page.set_viewport_size({'width':390,'height':844});assert not page.evaluate('document.documentElement.scrollWidth>innerWidth')
 capture(page.locator('[data-v14-audit-dashboard]'),'audit-results-mobile.png')
 dashboard();assert not page.evaluate('document.documentElement.scrollWidth>innerWidth')
 capture(page.locator('[data-v13-corrective-dashboard]'),'corrective-dashboard-mobile.png')
 tab('projects');page.get_by_role('button',name='ویرایش پروژه پروژه مشترک فعال',exact=True).click();form=page.get_by_role('dialog')
 form.locator('summary').filter(has_text='بودجه اقدام اصلاحی').click()
 assert '۳۵۰' in form.locator('[data-v12-budget-total]').inner_text() and '۶۵۰' in form.locator('[data-v12-budget-remaining]').inner_text()
 assert not page.evaluate('document.documentElement.scrollWidth>innerWidth')
 form.get_by_role('button',name='ذخیره پروژه',exact=True).click()
 assert not errors,errors
 assert all(request.split('#')[0]==args.url for request in requests),requests
 (art/'result.json').write_text(json.dumps({'result':'PASS','errors':errors,'requests':requests,'clauseDomains':8,'annexDomains':10,'attention':8,'projectCounts':{'total':5,'completed':3,'incomplete':2}},indent=2))
 print('PASS eight-domain reports, project budget/link persistence, JSON backup/restore, four-column automatic state after reload, dark/mobile/offline with no JavaScript errors or external requests.');browser.close()
